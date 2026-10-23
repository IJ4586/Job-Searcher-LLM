"""CareerCue backend: FastAPI + SQLite + a local LLM (Ollama) + job board search.

Everything is stored in one SQLite file on this machine. The only outbound
requests are to the job sources the user enables in Settings and to the local
Ollama server.
"""
from __future__ import annotations

import io
import json
import math
import os
import re
import sqlite3
import time
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

import httpx
from bs4 import BeautifulSoup
from fastapi import FastAPI, File, HTTPException, UploadFile
from pydantic import BaseModel, ConfigDict, Field, model_validator

BASE = Path(__file__).parent
DB_PATH = Path(os.getenv("CAREERCUE_DB", str(BASE / "careercue.db")))
DEFAULT_OLLAMA = os.getenv("OLLAMA_URL", "http://localhost:11434")
DEFAULT_MODEL = os.getenv("CAREERCUE_MODEL") or os.getenv("CLOSETCUE_MODEL") or "gemma4:e4b"
JOBBANK_SEARCH = "https://www.jobbank.gc.ca/jobsearch/jobsearch"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 CareerCue/1.0"}
BATCH = 5  # postings per LLM call; small local models stay accurate on short batches
STRONG = 70  # score that counts as a strong match in run summaries


# ---------- Database ----------
def conn():
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    return c


def rows(sql, args=()):
    with conn() as c:
        return [dict(r) for r in c.execute(sql, args).fetchall()]


def run(sql, args=()):
    with conn() as c:
        return c.execute(sql, args).lastrowid


def init_db():
    for table in ("experience", "preferences", "settings"):
        run(f"CREATE TABLE IF NOT EXISTS {table}(id INTEGER PRIMARY KEY, data TEXT)")
    run("""CREATE TABLE IF NOT EXISTS jobs(
        id INTEGER PRIMARY KEY,
        title TEXT, company TEXT, location TEXT, url TEXT, snippet TEXT,
        match_score INTEGER, match_reason TEXT,
        status TEXT DEFAULT 'new',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)""")
    have = {r["name"] for r in rows("PRAGMA table_info(jobs)")}
    extra = {
        "salary": "TEXT", "pay_value": "INTEGER", "posted": "TEXT", "posted_at": "TEXT",
        "source": "TEXT", "work_setting": "TEXT", "employment_type": "TEXT",
        "matched": "TEXT", "gap": "TEXT", "scored_by": "TEXT", "run_id": "INTEGER",
    }
    for col, kind in extra.items():
        if col not in have:
            run(f"ALTER TABLE jobs ADD COLUMN {col} {kind}")
    run("""CREATE TABLE IF NOT EXISTS runs(
        id INTEGER PRIMARY KEY,
        started_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        read INTEGER, strong INTEGER, scored_by TEXT, message TEXT)""")


# ---------- Data models ----------
class Doc(BaseModel):
    """Tolerant base: ignores unknown keys, treats null as "use the default",
    and accepts numbers where text is expected (local models do this a lot)."""
    model_config = ConfigDict(extra="ignore", coerce_numbers_to_str=True)

    @model_validator(mode="before")
    @classmethod
    def _drop_nulls(cls, v):
        return {k: x for k, x in v.items() if x is not None} if isinstance(v, dict) else v


class Basics(Doc):
    first_name: str = ""
    last_name: str = ""
    email: str = ""
    phone: str = ""
    career_stage: str = ""
    headline: str = ""
    country: str = ""
    region: str = ""
    city: str = ""
    linkedin: str = ""
    github: str = ""
    website: str = ""


class Work(Doc):
    title: str = ""
    company: str = ""
    location: str = ""
    employment_type: str = ""
    arrangement: str = ""
    start: str = ""
    end: str = ""
    current: bool = False
    description: str = ""
    skills: str = ""


class Project(Doc):
    name: str = ""
    role: str = ""
    link: str = ""
    start: str = ""
    end: str = ""
    tech: str = ""
    description: str = ""


class Education(Doc):
    institution: str = ""
    credential: str = ""
    field: str = ""
    location: str = ""
    start: str = ""
    end: str = ""
    current: bool = False
    grade: str = ""
    honours: str = ""
    coursework: str = ""


class Skills(Doc):
    technical: list[str] = Field(default_factory=list)
    tools: list[str] = Field(default_factory=list)
    soft: list[str] = Field(default_factory=list)


class Language(Doc):
    language: str = ""
    proficiency: str = ""


class Certification(Doc):
    name: str = ""
    issuer: str = ""
    issued: str = ""
    expires: str = ""
    credential: str = ""


class Activity(Doc):
    type: str = ""
    title: str = ""
    organization: str = ""
    location: str = ""
    date: str = ""
    details: str = ""


class Experience(Doc):
    basics: Basics = Field(default_factory=Basics)
    summary: str = ""
    work: list[Work] = Field(default_factory=list)
    projects: list[Project] = Field(default_factory=list)
    education: list[Education] = Field(default_factory=list)
    skills: Skills = Field(default_factory=Skills)
    languages: list[Language] = Field(default_factory=list)
    certifications: list[Certification] = Field(default_factory=list)
    activities: list[Activity] = Field(default_factory=list)


class Home(Doc):
    country: str = ""
    region: str = ""
    city: str = ""
    district: str = ""
    postal: str = ""


class Preferences(Doc):
    titles: list[str] = Field(default_factory=list, max_length=10)
    days: list[str] = Field(default_factory=lambda: ["Mon", "Tue", "Wed", "Thu", "Fri"])
    hours_start: str = "09:00"
    hours_end: str = "17:00"
    flexible_hours: bool = False
    shifts: list[str] = Field(default_factory=lambda: ["Day"])
    employment_types: list[str] = Field(default_factory=lambda: ["Full-time"])
    pay_basis: str = "annual"  # "annual" or "hourly"
    currency: str = "CAD"
    min_pay: float | None = None
    target_pay: float | None = None
    include_unlisted_pay: bool = True
    benefits: list[str] = Field(default_factory=list)
    work_settings: list[str] = Field(default_factory=list)  # empty = any
    work_areas: list[str] = Field(default_factory=list, max_length=3)
    max_commute: int = Field(45, ge=5, le=120)
    travel_mode: str = "Public transit"
    relocation: str = "Not open"
    work_authorization: str = ""
    available_from: str = ""
    notice_period: str = "Immediately"
    graduation: str = ""
    open_to_internships: bool = False
    start_immediately: bool = False
    home: Home = Field(default_factory=Home)
    posting_limit: int = Field(25, ge=1, le=500)
    skip_seen: bool = True

    @model_validator(mode="before")
    @classmethod
    def _nullable_pay(cls, v):
        # Keep explicit nulls for pay (they mean "not set"); the base drops them anyway.
        if isinstance(v, dict):
            for k in ("min_pay", "target_pay"):
                if v.get(k) == "":
                    v = {**v, k: None}
        return v


class Weights(Doc):
    skills: int = Field(30, ge=0, le=100)
    title: int = Field(25, ge=0, le=100)
    location: int = Field(20, ge=0, le=100)
    pay: int = Field(15, ge=0, le=100)
    level: int = Field(10, ge=0, le=100)


class Sources(Doc):
    jobbank: bool = True
    feeds: bool = False


class Settings(Doc):
    ollama_url: str = DEFAULT_OLLAMA
    model: str = ""  # empty = CAREERCUE_MODEL env var / built-in default
    temperature: float = Field(0.2, ge=0, le=2)
    num_ctx: int = Field(8192, ge=1024, le=131072)
    max_tokens: int = Field(1024, ge=128, le=8192)
    keep_loaded: bool = False
    strictness: str = "Balanced"  # Broad | Balanced | Strict
    weights: Weights = Field(default_factory=Weights)
    min_score: int = Field(0, ge=0, le=100)
    explain: bool = True
    gaps: bool = True
    sources: Sources = Field(default_factory=Sources)
    feed_urls: list[str] = Field(default_factory=list)


# ---------- Migrations from the first version's flat shapes ----------
def migrate_experience(d: dict) -> dict:
    if "basics" in d or not d:
        return d
    out: dict = {"summary": d.get("summary", "")}
    if d.get("skills"):
        out["skills"] = {"technical": d["skills"]}
    if d.get("history"):
        out["work"] = [{"description": d["history"]}]
    return out


def migrate_preferences(d: dict) -> dict:
    if "titles" in d or not d:
        return d
    out: dict = {"titles": d.get("roles", [])[:10]}
    if d.get("location"):
        out["work_areas"] = [d["location"]]
    remote = d.get("remote_pref", "")
    if remote in ("Remote", "Hybrid", "On-site"):
        out["work_settings"] = [remote]
    digits = re.sub(r"[^\d.]", "", d.get("min_salary", "") or "")
    if digits:
        try:
            out["min_pay"] = float(digits)
        except ValueError:
            pass
    return out


def load(table: str, model, migrate=None):
    r = rows(f"SELECT data FROM {table} WHERE id=1")
    data = json.loads(r[0]["data"]) if r else {}
    if migrate:
        data = migrate(data)
    return model.model_validate(data)


def save(table: str, doc: BaseModel):
    run(f"INSERT INTO {table}(id,data) VALUES(1,?) ON CONFLICT(id) DO UPDATE SET data=excluded.data",
        (doc.model_dump_json(),))
    return doc


def get_settings() -> Settings:
    return load("settings", Settings)


def model_name(s: Settings) -> str:
    return s.model.strip() or DEFAULT_MODEL


init_db()
app = FastAPI(title="CareerCue")


# ---------- Local LLM ----------
class LLMError(Exception):
    pass


def ask(s: Settings, prompt: str, timeout: float = 240, max_tokens: int | None = None):
    try:
        r = httpx.post(f"{s.ollama_url.rstrip('/')}/api/chat", timeout=timeout, json={
            "model": model_name(s),
            "messages": [{"role": "user", "content": prompt}],
            "stream": False,
            "format": "json",
            "keep_alive": "30m" if s.keep_loaded else "5m",
            "options": {"temperature": s.temperature, "num_ctx": s.num_ctx,
                        "num_predict": max_tokens or s.max_tokens},
        })
        r.raise_for_status()
        body = r.json()
        content = (body.get("message") or {}).get("content", "")
        if not content.strip():
            print("EMPTY REPLY FROM OLLAMA:", body)
            raise ValueError(f"The model returned an empty reply (done_reason={body.get('done_reason')})")
        return json.loads(content)
    except Exception as e:  # network, HTTP, or the model returned non-JSON
        raise LLMError(str(e)) from e


@app.get("/api/health")
def health():
    s = get_settings()
    name = model_name(s)
    try:
        res = httpx.get(f"{s.ollama_url.rstrip('/')}/api/tags", timeout=3)
        res.raise_for_status()
        names = [m["name"] for m in res.json().get("models", [])]
    except Exception as e:
        return {"ai": False, "reachable": False, "model": name, "models": [], "url": s.ollama_url,
                "error": f"Cannot reach Ollama at {s.ollama_url}: {e}"}
    installed = name in names or f"{name}:latest" in names
    return {"ai": installed, "reachable": True, "model": name, "models": names, "url": s.ollama_url,
            "error": "" if installed else f"Model '{name}' is not installed. Run: ollama pull {name}"}


@app.post("/api/settings/test")
def test_model():
    """Send a tiny prompt to prove the configured model loads and answers."""
    s = get_settings()
    t0 = time.time()
    try:
        out = ask(s, 'Reply with this JSON and nothing else: {"ok": true}', timeout=120)
    except LLMError as e:
        raise HTTPException(502, f"The model did not answer: {e}")
    return {"ok": bool(out.get("ok", True)), "model": model_name(s), "seconds": round(time.time() - t0, 1)}


# ---------- Experience ----------
@app.get("/api/experience")
def get_experience():
    return load("experience", Experience, migrate_experience)


@app.put("/api/experience")
def put_experience(e: Experience):
    return save("experience", e)


def extract_text(name: str, data: bytes) -> str:
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if ext == "pdf":
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in ("txt", "md"):
        return data.decode("utf-8", errors="ignore")
    raise HTTPException(415, "Upload a PDF, DOCX or TXT file.")


@app.post("/api/experience/import")
async def import_resume(file: UploadFile = File(...)):
    """Read a resume with the local model and return a draft. Nothing is saved
    until the user reviews the form and presses Save."""
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(413, "That file is over 5 MB.")
    text = extract_text(file.filename or "", data).strip()
    if len(text) < 40:
        raise HTTPException(422, "Could not read any text from that file. Scanned PDFs need OCR first.")
    schema = json.dumps(Experience(
        work=[Work()], projects=[Project()], education=[Education()],
        languages=[Language()], certifications=[Certification()], activities=[Activity()],
    ).model_dump(), indent=1)
    prompt = f"""Convert this resume into JSON with exactly the structure below.
Rules:
- Copy facts from the resume only. Never invent anything. Leave a field as "" when the resume does not say.
- Dates as YYYY-MM when the month is known, otherwise YYYY.
- "description" fields: one achievement per line, separated by newline characters.
- "skills" inside work entries and "tech" inside projects: comma-separated text.
- skills.technical, skills.tools and skills.soft are lists of short strings.
- Use empty lists for sections the resume does not have.
- career_stage is one of: Student, Recent graduate, Early career (0–3 years), Career changer, Returning to work, or "".

STRUCTURE:
{schema}

RESUME:
{text[:12000]}"""
    s = get_settings()
    try:
        out = ask(s, prompt, timeout=300, max_tokens=4096)
    except LLMError as e:
        raise HTTPException(502, f"The local model could not read the resume: {e}")
    return Experience.model_validate(out)


# ---------- Preferences ----------
@app.get("/api/preferences")
def get_preferences():
    return load("preferences", Preferences, migrate_preferences)


@app.put("/api/preferences")
def put_preferences(p: Preferences):
    return save("preferences", p)


# ---------- Settings ----------
@app.get("/api/settings")
def read_settings():
    return get_settings()


@app.put("/api/settings")
def put_settings(s: Settings):
    return save("settings", s)


# ---------- Job sources ----------
def parse_jobbank(html: str) -> list[dict]:
    out = []
    soup = BeautifulSoup(html, "html.parser")
    for article in soup.find_all("article", class_="action-buttons"):
        title_el = article.find("h3", class_="title")
        noc = title_el.find("span", class_="noctitle") if title_el else None
        title = noc.get_text(" ", strip=True) if noc else (title_el.get_text(" ", strip=True) if title_el else "")
        a_el = article.find("a", class_="resultJobItem")
        href = a_el.get("href", "") if a_el else ""
        link = ("https://www.jobbank.gc.ca" + href) if href.startswith("/") else href
        link = link.split(";jsessionid")[0]

        def li(cls):
            el = article.find("li", class_=cls)
            return el.get_text(" ", strip=True) if el else ""

        location = re.sub(r"^Location\s*", "", li("location")).strip()
        salary = re.sub(r"^Salary\s*", "", li("salary")).strip()
        posted = li("date")
        if not title:
            continue
        out.append({
            "title": title, "company": li("business"), "location": location, "url": link,
            "salary": salary, "posted": posted, "source": "Job Bank",
            "snippet": " ".join(x for x in [
                f"Employer: {li('business')}." if li("business") else "",
                f"Location: {location}." if location else "",
                f"Salary: {salary}." if salary else "",
                f"Posted: {posted}." if posted else "",
            ] if x),
        })
    return out


def search_jobbank(query: str, location: str, want: int, max_pages: int = 4) -> list[dict]:
    found: list[dict] = []
    for page in range(1, max_pages + 1):
        params = {"searchstring": query, "page": page, "sort": "M"}
        if location:
            params["locationstring"] = location
        try:
            resp = httpx.get(JOBBANK_SEARCH, params=params, headers=UA, timeout=20, follow_redirects=True)
        except httpx.HTTPError as e:
            print("Job Bank search error:", e)
            break
        if resp.status_code != 200:
            break
        batch = parse_jobbank(resp.text)
        if not batch:
            break
        found += batch
        if len(found) >= want:
            break
    return found


def search_feed(url: str) -> list[dict]:
    """Read an RSS or Atom feed of job postings."""
    try:
        resp = httpx.get(url, headers=UA, timeout=20, follow_redirects=True)
        resp.raise_for_status()
        root = ET.fromstring(resp.content)
    except Exception as e:
        print("Feed error:", url, e)
        return []
    atom = "{http://www.w3.org/2005/Atom}"
    out = []
    items = root.findall(".//item") or root.findall(f".//{atom}entry")
    host = re.sub(r"^https?://(www\.)?", "", url).split("/")[0]
    for it in items:
        def txt(*tags):
            for t in tags:
                el = it.find(t)
                if el is not None and (el.text or "").strip():
                    return el.text.strip()
            return ""
        link = txt("link")
        if not link:
            el = it.find(f"{atom}link")
            link = el.get("href", "") if el is not None else ""
        desc = BeautifulSoup(txt("description", f"{atom}summary", f"{atom}content"), "html.parser").get_text(" ", strip=True)
        title = txt("title", f"{atom}title")
        if title:
            out.append({"title": title, "company": "", "location": "", "url": link, "salary": "",
                        "posted": txt("pubDate", f"{atom}updated", f"{atom}published"),
                        "source": host, "snippet": desc[:1500]})
    return out


# ---------- Scoring ----------
def words(text: str) -> set[str]:
    return set(re.findall(r"[a-z0-9+#.]{2,}", text.lower()))


def all_skills(e: Experience) -> list[str]:
    s = list(e.skills.technical) + list(e.skills.tools) + list(e.skills.soft)
    for w in e.work:
        s += [x.strip() for x in w.skills.split(",") if x.strip()]
    for p in e.projects:
        s += [x.strip() for x in p.tech.split(",") if x.strip()]
    seen, out = set(), []
    for x in s:
        if x.lower() not in seen:
            seen.add(x.lower())
            out.append(x)
    return out


def profile_text(e: Experience) -> str:
    b = e.basics
    lines = []
    if b.career_stage or b.headline:
        lines.append(f"Stage: {b.career_stage}. Headline: {b.headline}".strip())
    if b.city or b.region or b.country:
        lines.append("Based in: " + ", ".join(x for x in [b.city, b.region, b.country] if x))
    if e.summary:
        lines.append("Summary: " + e.summary)
    for w in e.work:
        when = f"{w.start or '?'} to {'present' if w.current else (w.end or '?')}"
        lines.append(f"Work: {w.title} at {w.company} ({when}; {w.employment_type or 'type not given'}). "
                     f"{w.description[:500]} Skills: {w.skills}")
    for p in e.projects:
        lines.append(f"Project: {p.name} ({p.role}). Tech: {p.tech}. {p.description[:300]}")
    for ed in e.education:
        when = "in progress" if ed.current else (ed.end or "")
        lines.append(f"Education: {ed.credential} in {ed.field}, {ed.institution} ({when}). Coursework: {ed.coursework}")
    skills = all_skills(e)
    if skills:
        lines.append("Skills: " + ", ".join(skills))
    if e.languages:
        lines.append("Languages: " + ", ".join(f"{l.language} ({l.proficiency})" for l in e.languages if l.language))
    for c in e.certifications:
        lines.append(f"Certification: {c.name} ({c.issuer})")
    for a in e.activities:
        lines.append(f"{a.type or 'Activity'}: {a.title}, {a.organization}. {a.details[:200]}")
    return "\n".join(lines)[:7000] or "No experience entered yet."


def prefs_text(p: Preferences) -> str:
    unit = "per hour" if p.pay_basis == "hourly" else "per year"
    pay = []
    if p.min_pay:
        pay.append(f"minimum {p.currency} {p.min_pay:g} {unit}")
    if p.target_pay:
        pay.append(f"target {p.currency} {p.target_pay:g} {unit}")
    home = ", ".join(x for x in [p.home.district, p.home.city, p.home.region, p.home.country] if x)
    lines = [
        "Target titles: " + (", ".join(p.titles) or "any"),
        "Employment types: " + (", ".join(p.employment_types) or "any"),
        f"Schedule: {', '.join(p.days) or 'any days'}, {p.hours_start}–{p.hours_end}"
        + (", flexible hours fine" if p.flexible_hours else "") + f". Shifts: {', '.join(p.shifts) or 'any'}",
        "Pay: " + ("; ".join(pay) or "not specified")
        + ("" if p.include_unlisted_pay else ". Postings without pay are less welcome"),
        "Work setting: " + (", ".join(p.work_settings) or "any"),
        "Preferred work areas: " + (", ".join(p.work_areas) or "anywhere nearby"),
        f"Lives in: {home or 'not given'}. Max commute {p.max_commute} min by {p.travel_mode.lower()}. "
        f"Relocation: {p.relocation}",
        f"Available from: {p.available_from or 'now'}; notice: {p.notice_period}"
        + (f"; graduating {p.graduation}" if p.graduation else "")
        + ("; open to internships/co-op" if p.open_to_internships else ""),
    ]
    if p.work_authorization:
        lines.append(f"Work authorization: {p.work_authorization}")
    if p.benefits:
        lines.append("Valued benefits: " + ", ".join(p.benefits))
    return "\n".join(lines)


STRICTNESS = {
    "Broad": "Be generous. Give partial credit for transferable skills and adjacent titles.",
    "Balanced": "Be fair. Close matches score high, stretch roles score in the middle.",
    "Strict": "Be strict. Score above 70 only when the posting meets nearly every stated requirement and preference.",
}


def score_prompt(batch: list[dict], profile: str, prefs: str, s: Settings) -> str:
    w = s.weights
    total = (w.skills + w.title + w.location + w.pay + w.level) or 1
    pct = lambda x: round(100 * x / total)
    postings = [{"id": i, **{k: j[k] for k in ("title", "company", "location", "salary", "posted")},
                 "text": j["snippet"][:1500]} for i, j in enumerate(batch)]
    reason = ('"reason": "one or two plain sentences to the candidate (use \\"you\\") on why it fits or not"'
              if s.explain else '"reason": ""')
    gap = ('"gap": "the single biggest requirement or preference this posting misses, or \\"\\" if none"'
           if s.gaps else '"gap": ""')
    return f"""You score job postings for one candidate. For each posting give a score from 0 to 100 for how well it fits the candidate's experience AND preferences.

Weighting (share of the score):
- Skills match: {pct(w.skills)}%
- Job title fit: {pct(w.title)}%
- Location, setting and commute: {pct(w.location)}%
- Pay: {pct(w.pay)}%
- Experience level fit: {pct(w.level)}%
{STRICTNESS.get(s.strictness, STRICTNESS['Balanced'])}
Use only facts in the posting and the candidate profile. Never invent requirements. If the posting does not say something, do not penalise heavily for it.

CANDIDATE
{profile}

PREFERENCES
{prefs}

POSTINGS
{json.dumps(postings, ensure_ascii=False)}

Reply with JSON only, one result per posting id:
{{"results": [{{"id": 0, "score": 0, {reason}, "matched": ["up to 4 of the candidate's skills or facts the posting asks for"], {gap}, "setting": "Remote|Hybrid|On-site|Unknown", "employment_type": "Full-time|Part-time|Contract|Temporary|Internship|Co-op|Unknown"}}]}}"""


def keyword_score(job: dict, e: Experience, p: Preferences) -> dict:
    """Fallback when the model is unavailable: transparent keyword overlap."""
    text = f"{job['title']} {job['snippet']}".lower()
    title_words = words(job["title"])
    best_title = max((len(words(t) & title_words) / max(1, len(words(t))) for t in p.titles), default=0)
    hits = [s for s in all_skills(e) if s.lower() in text]
    score = min(100, round(30 + 45 * best_title + 6 * len(hits)))
    return {"score": score, "matched": hits[:4], "gap": "", "setting": "Unknown", "employment_type": "Unknown",
            "reason": "Scored by keyword overlap only, because the local model was not available."}


def score_jobs(jobs: list[dict], e: Experience, p: Preferences, s: Settings) -> tuple[list[dict], str]:
    profile, prefs = profile_text(e), prefs_text(p)
    used = set()
    for i in range(0, len(jobs), BATCH):
        batch = jobs[i:i + BATCH]
        results: dict[int, dict] = {}
        try:
            out = ask(s, score_prompt(batch, profile, prefs, s))
            for r in out.get("results") or out.get("jobs") or []:
                if isinstance(r, dict) and str(r.get("id", "")).isdigit():
                    results[int(r["id"])] = r
            used.add("llm")
        except LLMError as err:
            print("LLM scoring failed, using keywords:", err)
        for k, job in enumerate(batch):
            r = results.get(k)
            if r is None:
                r = keyword_score(job, e, p)
                used.add("keywords")
                job["scored_by"] = "keywords"
            else:
                job["scored_by"] = "llm"
            try:
                job["match_score"] = max(0, min(100, int(float(r.get("score", 0)))))
            except (TypeError, ValueError):
                job["match_score"] = 0
            job["match_reason"] = str(r.get("reason") or "")
            matched = r.get("matched") or []
            job["matched"] = [str(m) for m in matched if m][:4] if isinstance(matched, list) else []
            job["gap"] = str(r.get("gap") or "")
            setting = str(r.get("setting") or "Unknown")
            job["work_setting"] = setting if setting in ("Remote", "Hybrid", "On-site") else ""
            etype = str(r.get("employment_type") or "Unknown")
            job["employment_type"] = "" if etype == "Unknown" else etype
    by = "mixed" if len(used) > 1 else (used.pop() if used else "keywords")
    return jobs, by


def annual_pay(salary: str) -> int | None:
    nums = [float(n.replace(",", "")) for n in re.findall(r"\d[\d,]*(?:\.\d+)?", salary or "")]
    if not nums:
        return None
    v = max(nums)
    low = salary.lower()
    if "hour" in low or v < 300:
        v *= 2080
    elif "week" in low:
        v *= 52
    elif "month" in low:
        v *= 12
    return int(v)


def posted_iso(text: str) -> str | None:
    clean = re.sub(r"^(Posted|Date posted)[:\s]*", "", (text or "").strip(), flags=re.I)
    for fmt in ("%B %d, %Y", "%b %d, %Y", "%Y-%m-%d", "%a, %d %b %Y %H:%M:%S %z", "%a, %d %b %Y %H:%M:%S %Z"):
        try:
            return datetime.strptime(clean, fmt).date().isoformat()
        except ValueError:
            continue
    m = re.match(r"(\d{4}-\d{2}-\d{2})", clean)
    return m.group(1) if m else None


# ---------- Jobs ----------
def job_out(r: dict) -> dict:
    r = dict(r)
    try:
        r["matched"] = json.loads(r.get("matched") or "[]")
    except ValueError:
        r["matched"] = []
    return r


@app.post("/api/jobs/search")
def search_jobs():
    e = get_experience()
    p = get_preferences()
    s = get_settings()
    if not p.titles:
        raise HTTPException(400, "Add at least one target job title in Job preferences first.")
    if not (s.sources.jobbank or (s.sources.feeds and any(u.strip() for u in s.feed_urls))):
        raise HTTPException(400, "Turn on at least one job source in Settings.")

    limit = p.posting_limit
    seen = {r["url"] for r in rows("SELECT url FROM jobs WHERE url != ''")} if p.skip_seen else set()
    remote_only = p.work_settings == ["Remote"]
    location = "" if remote_only else (", ".join(x for x in [p.home.city, p.home.region] if x)
                                       or (p.work_areas[0] if p.work_areas else ""))

    collected: list[dict] = []
    keys: set[str] = set()

    def add(items: list[dict], cap: int) -> None:
        for j in items:
            key = j["url"] or f"{j['title']}|{j['company']}".lower()
            if key in keys or (j["url"] and j["url"] in seen):
                continue
            keys.add(key)
            collected.append(j)
            if len(collected) >= cap:
                return

    if s.sources.jobbank:
        per_title = math.ceil(limit / len(p.titles))
        for i, title in enumerate(p.titles):
            cap = min(limit, per_title * (i + 1))
            # Ask for extra when skipping seen postings, since some results will be dropped.
            add(search_jobbank(title, location, want=per_title * (2 if seen else 1)), cap)
    if s.sources.feeds and len(collected) < limit:
        targets = set().union(*(words(t) for t in p.titles))
        feed_items: list[dict] = []
        for url in s.feed_urls:
            if url.strip():
                feed_items += search_feed(url.strip())
        feed_items.sort(key=lambda j: -len(words(j["title"]) & targets))
        add(feed_items, limit)

    if not collected:
        run("INSERT INTO runs(read, strong, scored_by, message) VALUES(0,0,'',?)",
            ("No new postings found. Try broader titles, a different location, or turn off 'skip seen'.",))
        return {"ok": True, "read": 0, "strong": 0, "scored_by": "", "jobs": [],
                "message": "No new postings found. Try broader titles, a different location, or turn off 'skip seen'."}

    scored, by = score_jobs(collected[:limit], e, p, s)
    run_id = run("INSERT INTO runs(read, strong, scored_by, message) VALUES(?,?,?,'')",
                 (len(scored), sum(1 for j in scored if j["match_score"] >= STRONG), by))
    for j in scored:
        run("""INSERT INTO jobs(title, company, location, url, snippet, match_score, match_reason, salary, pay_value,
                                posted, posted_at, source, work_setting, employment_type, matched, gap, scored_by, run_id)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (j["title"], j["company"], j["location"], j["url"], j["snippet"], j["match_score"], j["match_reason"],
             j["salary"], annual_pay(j["salary"]), j["posted"], posted_iso(j["posted"]), j["source"],
             j["work_setting"], j["employment_type"], json.dumps(j["matched"]), j["gap"], j["scored_by"], run_id))
    return {"ok": True, "read": len(scored), "strong": sum(1 for j in scored if j["match_score"] >= STRONG),
            "scored_by": by, "message": "", "jobs": list_jobs()}


@app.get("/api/jobs")
def list_jobs(include_hidden: bool = False):
    where = "" if include_hidden else "WHERE status != 'hidden'"
    return [job_out(r) for r in rows(f"SELECT * FROM jobs {where} ORDER BY match_score DESC, id DESC")]


@app.get("/api/runs/last")
def last_run():
    r = rows("SELECT * FROM runs ORDER BY id DESC LIMIT 1")
    return r[0] if r else None


class StatusIn(BaseModel):
    status: str = Field(pattern="^(new|saved|hidden)$")


@app.patch("/api/jobs/{job_id}")
def set_job_status(job_id: int, body: StatusIn):
    run("UPDATE jobs SET status=? WHERE id=?", (body.status, job_id))
    r = rows("SELECT * FROM jobs WHERE id=?", (job_id,))
    if not r:
        raise HTTPException(404, "Posting not found.")
    return job_out(r[0])


@app.delete("/api/jobs/{job_id}")
def delete_job(job_id: int):
    run("DELETE FROM jobs WHERE id=?", (job_id,))
    return {"ok": True}


@app.delete("/api/jobs")
def clear_jobs(keep_saved: bool = True):
    run("DELETE FROM jobs" + (" WHERE status != 'saved'" if keep_saved else ""))
    return {"ok": True}


# ---------- Your data ----------
@app.get("/api/export")
def export_all():
    return {"exported_at": datetime.now().isoformat(timespec="seconds"),
            "experience": get_experience().model_dump(), "preferences": get_preferences().model_dump(),
            "settings": get_settings().model_dump(), "jobs": list_jobs(include_hidden=True)}


@app.delete("/api/data")
def delete_all():
    for t in ("experience", "preferences", "settings", "jobs", "runs"):
        run(f"DELETE FROM {t}")
    return {"ok": True}